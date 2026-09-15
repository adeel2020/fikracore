#!/usr/bin/env python3
"""
Generates a 100% standard-compliant, beautifully designed PowerPoint presentation (.pptx)
detailing all use cases implemented using gbrain in the kagent / MARK platform.
Uses python-pptx for guaranteed compatibility with Microsoft PowerPoint, Apple Keynote,
Google Slides, and LibreOffice.
"""

import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

# Slide Dimensions (16:9 Widescreen)
SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)

# Color Palette
NAVY_DARK = RGBColor(15, 23, 42)       # Slate 900 #0F172A
NAVY_CARD = RGBColor(30, 41, 59)      # Slate 800 #1E293B
SLATE_BG = RGBColor(248, 250, 252)    # Slate 50  #F8FAFC
CARD_BG = RGBColor(255, 255, 255)     # Pure White
BORDER_LIGHT = RGBColor(226, 232, 240)# Slate 200 #E2E8F0
BORDER_DARK = RGBColor(51, 65, 85)    # Slate 700 #334155

TEXT_DARK = RGBColor(15, 23, 42)      # Slate 900
TEXT_MUTED = RGBColor(71, 85, 105)    # Slate 600
TEXT_LIGHT = RGBColor(148, 163, 184)  # Slate 400
WHITE = RGBColor(255, 255, 255)

BLUE_ACCENT = RGBColor(37, 99, 235)   # Blue 600   #2563EB
BLUE_LIGHT = RGBColor(219, 234, 254)  # Blue 100
CYAN_ACCENT = RGBColor(2, 132, 199)   # Sky 600    #0284C7
EMERALD = RGBColor(5, 150, 105)       # Emerald 600#059669
AMBER = RGBColor(217, 119, 6)         # Amber 600  #D97706
PURPLE = RGBColor(124, 58, 237)       # Violet 600 #7C3AED
RED = RGBColor(220, 38, 38)           # Red 600    #DC2626


def set_bg(slide, color):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_header(slide, badge: str, title: str, subtitle: str = "", is_dark: bool = False):
    # Category badge
    badge_box = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.4), Inches(2.6), Inches(0.32)
    )
    badge_box.fill.solid()
    badge_box.fill.fore_color.rgb = BLUE_ACCENT if not is_dark else BLUE_LIGHT
    badge_box.line.fill.background()
    tf_b = badge_box.text_frame
    tf_b.word_wrap = False
    p_b = tf_b.paragraphs[0]
    p_b.alignment = PP_ALIGN.CENTER
    run_b = p_b.add_run()
    run_b.text = badge.upper()
    run_b.font.name = "Segoe UI"
    run_b.font.size = Pt(10)
    run_b.font.bold = True
    run_b.font.color.rgb = WHITE if not is_dark else NAVY_DARK

    # Title & Subtitle box
    t_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.78), Inches(11.7), Inches(0.95))
    tf_t = t_box.text_frame
    tf_t.word_wrap = True
    p_t = tf_t.paragraphs[0]
    run_t = p_t.add_run()
    run_t.text = title
    run_t.font.name = "Segoe UI"
    run_t.font.size = Pt(22)
    run_t.font.bold = True
    run_t.font.color.rgb = TEXT_DARK if not is_dark else WHITE

    if subtitle:
        p_sub = tf_t.add_paragraph()
        p_sub.space_before = Pt(3)
        run_sub = p_sub.add_run()
        run_sub.text = subtitle
        run_sub.font.name = "Segoe UI"
        run_sub.font.size = Pt(12)
        run_sub.font.color.rgb = TEXT_MUTED if not is_dark else TEXT_LIGHT

    # Footer
    f_box = slide.shapes.add_textbox(Inches(0.8), Inches(7.05), Inches(11.7), Inches(0.3))
    tf_f = f_box.text_frame
    p_f = tf_f.paragraphs[0]
    run_f = p_f.add_run()
    run_f.text = "Autonomous Telecom Operations with gbrain | KAgent MARK Architecture"
    run_f.font.name = "Segoe UI"
    run_f.font.size = Pt(9)
    run_f.font.color.rgb = TEXT_LIGHT if not is_dark else TEXT_MUTED


def add_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=BORDER_LIGHT):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    if border_color:
        card.line.color.rgb = border_color
        card.line.width = Pt(1)
    else:
        card.line.fill.background()
    return card


def build_all_slides(prs):
    blank_layout = prs.slide_layouts[6]

    # =========================================================================
    # SLIDE 1: Title Slide (Dark Theme)
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_bg(s1, NAVY_DARK)
    card1 = add_card(s1, Inches(0.8), Inches(0.8), Inches(11.73), Inches(5.9), bg_color=NAVY_CARD, border_color=BORDER_DARK)

    # Badge
    b1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.4), Inches(1.4), Inches(3.2), Inches(0.38))
    b1.fill.solid()
    b1.fill.fore_color.rgb = BLUE_ACCENT
    b1.line.fill.background()
    p_b1 = b1.text_frame.paragraphs[0]
    p_b1.alignment = PP_ALIGN.CENTER
    r_b1 = p_b1.add_run()
    r_b1.text = "ENTERPRISE TELECOM AIOPS"
    r_b1.font.name = "Segoe UI"
    r_b1.font.size = Pt(11)
    r_b1.font.bold = True
    r_b1.font.color.rgb = WHITE

    # Title box
    tb1 = s1.shapes.add_textbox(Inches(1.4), Inches(2.0), Inches(10.5), Inches(4.2))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p1 = tf1.paragraphs[0]
    r1 = p1.add_run()
    r1.text = "Autonomous Telecom Operations\nwith gbrain"
    r1.font.name = "Segoe UI"
    r1.font.size = Pt(36)
    r1.font.bold = True
    r1.font.color.rgb = WHITE

    p2 = tf1.add_paragraph()
    p2.space_before = Pt(14)
    r2 = p2.add_run()
    r2.text = "Semantic Knowledge Graphs, Multi-Domain Correlation & Deterministic Storytelling"
    r2.font.name = "Segoe UI"
    r2.font.size = Pt(17)
    r2.font.color.rgb = CYAN_ACCENT

    p3 = tf1.add_paragraph()
    p3.space_before = Pt(16)
    r3 = p3.add_run()
    r3.text = "Architecture, Integration Patterns & Production-Grade Implemented Use Cases"
    r3.font.name = "Segoe UI"
    r3.font.size = Pt(13)
    r3.font.color.rgb = TEXT_LIGHT

    p4 = tf1.add_paragraph()
    p4.space_before = Pt(28)
    r4 = p4.add_run()
    r4.text = "KAgent Platform  |  MARK Intelligent Telecom Assistant  |  2026 Production Release"
    r4.font.name = "Segoe UI"
    r4.font.size = Pt(12)
    r4.font.bold = True
    r4.font.color.rgb = EMERALD

    # =========================================================================
    # SLIDE 2: Executive Summary & The Telecom Problem
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_bg(s2, SLATE_BG)
    add_header(s2, "Executive Summary", "The Challenge of Modern Telecom NOC Operations", "Why traditional alerting fails and how gbrain transforms cognitive incident management")

    card_w = Inches(3.64)
    card_h = Inches(4.9)
    card_y = Inches(1.85)

    # Card 1: Crisis
    add_card(s2, Inches(0.8), card_y, card_w, card_h)
    tb = s2.shapes.add_textbox(Inches(0.9), card_y + Inches(0.15), card_w - Inches(0.2), card_h - Inches(0.3))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "The NOC Operational Crisis"
    r.font.bold = True
    r.font.size = Pt(15)
    r.font.color.rgb = RED

    items1 = [
        ("Alarm Storms: ", "Over 100k alarms/day across RAN, Core, and Transport silos creating severe alert fatigue."),
        ("Fragmented Visibility: ", "Disjointed tools for metrics, logs, tickets, and topology."),
        ("Tribal Knowledge: ", "RCA expertise locked in senior engineers' heads rather than machine-executable models."),
        ("High MTTR: ", "Hours lost triaging symptoms instead of remediating root causes."),
    ]
    for b_title, b_text in items1:
        p_item = tf.add_paragraph()
        p_item.space_before = Pt(10)
        p_item.level = 0
        r_bt = p_item.add_run()
        r_bt.text = "• " + b_title
        r_bt.font.bold = True
        r_bt.font.size = Pt(11)
        r_bt.font.color.rgb = TEXT_DARK
        r_bc = p_item.add_run()
        r_bc.text = b_text
        r_bc.font.size = Pt(11)
        r_bc.font.color.rgb = TEXT_MUTED

    # Card 2: Generic LLM Flaw
    add_card(s2, Inches(4.84), card_y, card_w, card_h)
    tb = s2.shapes.add_textbox(Inches(4.94), card_y + Inches(0.15), card_w - Inches(0.2), card_h - Inches(0.3))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "The Flaw of Generic LLMs"
    r.font.bold = True
    r.font.size = Pt(15)
    r.font.color.rgb = AMBER

    items2 = [
        ("Hallucination Risk: ", "Inventing non-existent network functions, wrong alarm causes, or fictitious timestamps."),
        ("Context Degradation: ", "Dumping raw telemetry tokens into prompts exhausts context and lowers precision."),
        ("Non-Deterministic: ", "Produces different causal conclusions for identical production outages."),
        ("Missing Provenance: ", "Cannot trace reasoning steps back to exact Prometheus, Loki, or NMS records."),
    ]
    for b_title, b_text in items2:
        p_item = tf.add_paragraph()
        p_item.space_before = Pt(10)
        p_item.level = 0
        r_bt = p_item.add_run()
        r_bt.text = "• " + b_title
        r_bt.font.bold = True
        r_bt.font.size = Pt(11)
        r_bt.font.color.rgb = TEXT_DARK
        r_bc = p_item.add_run()
        r_bc.text = b_text
        r_bc.font.size = Pt(11)
        r_bc.font.color.rgb = TEXT_MUTED

    # Card 3: The gbrain Paradigm
    add_card(s2, Inches(8.88), card_y, card_w, card_h)
    tb = s2.shapes.add_textbox(Inches(8.98), card_y + Inches(0.15), card_w - Inches(0.2), card_h - Inches(0.3))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "The gbrain Paradigm"
    r.font.bold = True
    r.font.size = Pt(15)
    r.font.color.rgb = EMERALD

    items3 = [
        ("Semantic Graph Hub: ", "Maintains explicit graph topology of functions, symptoms, and causal paths."),
        ("Deterministic Narratives: ", "Verifiable storytelling backed by structured knowledge graph traversal."),
        ("MCP Interoperability: ", "Standardized tool contracts exposed via Model Context Protocol (MCP)."),
        ("Closed-Loop Learning: ", "FCAPS operational lens continuously updates runbooks and query templates."),
    ]
    for b_title, b_text in items3:
        p_item = tf.add_paragraph()
        p_item.space_before = Pt(10)
        p_item.level = 0
        r_bt = p_item.add_run()
        r_bt.text = "• " + b_title
        r_bt.font.bold = True
        r_bt.font.size = Pt(11)
        r_bt.font.color.rgb = TEXT_DARK
        r_bc = p_item.add_run()
        r_bc.text = b_text
        r_bc.font.size = Pt(11)
        r_bc.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 3: Architecture & Core Design Principle
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_bg(s3, SLATE_BG)
    add_header(s3, "Platform Architecture", "Core Principle: Semantic Brain vs. Raw Data Store", "Clean separation of concerns guarantees zero data duplication and instant graph scalability")

    # Banner
    add_card(s3, Inches(0.8), Inches(1.85), Inches(11.7), Inches(0.95), bg_color=RGBColor(239, 246, 255), border_color=RGBColor(147, 197, 253))
    tb_ban = s3.shapes.add_textbox(Inches(1.0), Inches(1.9), Inches(11.3), Inches(0.85))
    tf_ban = tb_ban.text_frame
    tf_ban.word_wrap = True
    p_ban1 = tf_ban.paragraphs[0]
    r_ban1a = p_ban1.add_run()
    r_ban1a.text = "Architectural Golden Rule: "
    r_ban1a.font.bold = True
    r_ban1a.font.size = Pt(13)
    r_ban1a.font.color.rgb = BLUE_ACCENT
    r_ban1b = p_ban1.add_run()
    r_ban1b.text = "\"gbrain is the semantic telecom brain, NOT the raw source-of-truth store.\""
    r_ban1b.font.bold = True
    r_ban1b.font.size = Pt(13)
    r_ban1b.font.color.rgb = NAVY_DARK

    p_ban2 = tf_ban.add_paragraph()
    p_ban2.space_before = Pt(3)
    r_ban2 = p_ban2.add_run()
    r_ban2.text = "Raw incidents stay in Incident Registry. Raw telemetry stays in Grafana/LGTM. Raw alarms stay in NMS. gbrain stores compact semantic projections, causal links, and domain ontologies."
    r_ban2.font.size = Pt(10.5)
    r_ban2.font.color.rgb = TEXT_MUTED

    # 2 Comparison Columns
    col_w = Inches(5.7)
    col_h = Inches(3.8)
    col_y = Inches(3.0)

    # Left: Raw Sources
    add_card(s3, Inches(0.8), col_y, col_w, col_h)
    tb = s3.shapes.add_textbox(Inches(1.0), col_y + Inches(0.15), col_w - Inches(0.4), col_h - Inches(0.3))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "Raw Operational Layer (Systems of Record)"
    r.font.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = NAVY_DARK

    col1_items = [
        ("Grafana / LGTM Stack: ", "High-velocity metrics (Mimir), logs (Loki), traces (Tempo)."),
        ("NMS / EMS / OSS: ", "Raw alarm feeds from Nokia NetAct, Ericsson ENM, Huawei U2000."),
        ("ITSM & Incident Registry: ", "Remedy / ServiceNow tickets, SLA timers, operator queues."),
        ("Data Nature: ", "Petabyte-scale, high-frequency, unlinked, siloed schemas."),
    ]
    for bt, bc in col1_items:
        pi = tf.add_paragraph()
        pi.space_before = Pt(10)
        r1 = pi.add_run()
        r1.text = "• " + bt
        r1.font.bold = True
        r1.font.size = Pt(11)
        r1.font.color.rgb = TEXT_DARK
        r2 = pi.add_run()
        r2.text = bc
        r2.font.size = Pt(11)
        r2.font.color.rgb = TEXT_MUTED

    # Right: gbrain
    add_card(s3, Inches(6.8), col_y, col_w, col_h)
    tb = s3.shapes.add_textbox(Inches(7.0), col_y + Inches(0.15), col_w - Inches(0.4), col_h - Inches(0.3))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "gbrain Semantic Brain (Cognitive Layer)"
    r.font.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = BLUE_ACCENT

    col2_items = [
        ("Compact Projections: ", "Stores semantic summaries (e.g. 'CPU @ 98%', 'RSR breach')."),
        ("Causal Graph Edges: ", "Explicit links (affects, has-hypothesis, supported-by, targets)."),
        ("Domain Knowledge: ", "3GPP functional models, standard procedures, service intents."),
        ("Deterministic Access: ", "Zero hallucination: only traversed facts enter the narrative."),
    ]
    for bt, bc in col2_items:
        pi = tf.add_paragraph()
        pi.space_before = Pt(10)
        r1 = pi.add_run()
        r1.text = "• " + bt
        r1.font.bold = True
        r1.font.size = Pt(11)
        r1.font.color.rgb = TEXT_DARK
        r2 = pi.add_run()
        r2.text = bc
        r2.font.size = Pt(11)
        r2.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 4: MCP Hub & Transport Architecture
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_bg(s4, SLATE_BG)
    add_header(s4, "System Integration", "Model Context Protocol (MCP) Hub & Transport", "Standardized outbound client hub connecting MARK orchestrator to gbrain and external tools")

    add_card(s4, Inches(0.8), Inches(1.85), col_w, Inches(4.9))
    tb = s4.shapes.add_textbox(Inches(1.0), Inches(2.0), col_w - Inches(0.4), Inches(4.6))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "Centralized MCP Client Hub"
    r.font.bold = True
    r.font.size = Pt(15)
    r.font.color.rgb = NAVY_DARK

    mcp_items = [
        ("Single Outbound Gateway: ", "All services (Storyteller, Correlation, RCA) call gbrain through `mcp_client_hub`."),
        ("No One-Off Clients: ", "Prevents ad-hoc socket connections and unmonitored API calls across services."),
        ("Connector Manifests: ", "YAML definitions governing permissions, rate limits, and audit policies."),
        ("End-to-End Tracing: ", "Every tool call produces structured JSON traces for compliance and auditing."),
        ("Enterprise Tooling: ", "Exposes status & health via GET /api/jarvis/connectors/status."),
    ]
    for bt, bc in mcp_items:
        pi = tf.add_paragraph()
        pi.space_before = Pt(10)
        r1 = pi.add_run()
        r1.text = "• " + bt
        r1.font.bold = True
        r1.font.size = Pt(11)
        r1.font.color.rgb = TEXT_DARK
        r2 = pi.add_run()
        r2.text = bc
        r2.font.size = Pt(11)
        r2.font.color.rgb = TEXT_MUTED

    add_card(s4, Inches(6.8), Inches(1.85), col_w, Inches(4.9))
    tb = s4.shapes.add_textbox(Inches(7.0), Inches(2.0), col_w - Inches(0.4), Inches(4.6))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "Four-Tier Resilient Transport"
    r.font.bold = True
    r.font.size = Pt(15)
    r.font.color.rgb = BLUE_ACCENT

    trans_items = [
        ("1. HTTP MCP (Primary): ", "JSON-RPC 2.0 over HTTP (GBRAIN_MCP_URL=http://localhost:3131/mcp) with bearer token auth."),
        ("2. Persistent Stdio MCP: ", "Process-bound stdio JSON lines for containerized environments."),
        ("3. Subprocess CLI Fallback: ", "Direct execution via `gbrain call <tool>` when daemon process is offline."),
        ("4. Embedded In-Memory Graph: ", "Built-in mobile core fixture catalog ensuring zero downtime during network partitioning."),
    ]
    for bt, bc in trans_items:
        pi = tf.add_paragraph()
        pi.space_before = Pt(10)
        r1 = pi.add_run()
        r1.text = "• " + bt
        r1.font.bold = True
        r1.font.size = Pt(11)
        r1.font.color.rgb = TEXT_DARK
        r2 = pi.add_run()
        r2.text = bc
        r2.font.size = Pt(11)
        r2.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 5: gbrain Knowledge Model & Namespace Hierarchy
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_bg(s5, SLATE_BG)
    add_header(s5, "Knowledge Taxonomy", "gbrain Namespace Hierarchy & Entity Model", "Structured domain partitioning prevents data pollution and maintains strict ontological boundaries")

    gw = Inches(5.7)
    gh = Inches(2.35)

    # Box 1: knowledge/
    add_card(s5, Inches(0.8), Inches(1.85), gw, gh)
    tb = s5.shapes.add_textbox(Inches(1.0), Inches(1.95), gw - Inches(0.4), gh - Inches(0.2))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "knowledge/mobile-core/... (Stable Domain Knowledge)"
    r.font.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = NAVY_DARK
    k_items = [
        ("3GPP Architecture: ", "AMF, SMF, UPF, gNodeB, MME, HSS, P-CSCF specifications."),
        ("Standard Topology: ", "Interface connections (N1, N2, N3, S1-MME, SGi, Diameter)."),
        ("Service Procedures: ", "Initial UE Registration, PDU Session Establishment, VoLTE Call."),
    ]
    for bt, bc in k_items:
        pi = tf.add_paragraph()
        pi.space_before = Pt(6)
        r1 = pi.add_run()
        r1.text = "• " + bt
        r1.font.bold = True
        r1.font.size = Pt(10)
        r1.font.color.rgb = TEXT_DARK
        r2 = pi.add_run()
        r2.text = bc
        r2.font.size = Pt(10)
        r2.font.color.rgb = TEXT_MUTED

    # Box 2: observations/
    add_card(s5, Inches(6.8), Inches(1.85), gw, gh)
    tb = s5.shapes.add_textbox(Inches(7.0), Inches(1.95), gw - Inches(0.4), gh - Inches(0.2))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "observations/mobile-core/... (Semantic Facts)"
    r.font.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = BLUE_ACCENT
    o_items = [
        ("Projected Alarms: ", "Normalized alarms with domain, severity, and object ID."),
        ("KPI Breaches: ", "RSR dropped to 94.7%, 4G_Attach_SR degraded to 88.2%."),
        ("Telemetry Evidence: ", "AMF-01 CPU 98%, S1AP message drops, SIP 503 timeouts."),
    ]
    for bt, bc in o_items:
        pi = tf.add_paragraph()
        pi.space_before = Pt(6)
        r1 = pi.add_run()
        r1.text = "• " + bt
        r1.font.bold = True
        r1.font.size = Pt(10)
        r1.font.color.rgb = TEXT_DARK
        r2 = pi.add_run()
        r2.text = bc
        r2.font.size = Pt(10)
        r2.font.color.rgb = TEXT_MUTED

    # Box 3: correlation/ & incidents/
    add_card(s5, Inches(0.8), Inches(4.45), gw, gh)
    tb = s5.shapes.add_textbox(Inches(1.0), Inches(4.55), gw - Inches(0.4), gh - Inches(0.2))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "correlation/ & incidents/... (Reasoning Entities)"
    r.font.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = PURPLE
    c_items = [
        ("Correlation Clusters: ", "Multi-alarm grouping by time, topology, and reachability."),
        ("Hypotheses: ", "Plausible causal assertions with confidence scores (0.0 - 1.0)."),
        ("Decisions & Incidents: ", "Promotion decisions and canonical incident pages."),
    ]
    for bt, bc in c_items:
        pi = tf.add_paragraph()
        pi.space_before = Pt(6)
        r1 = pi.add_run()
        r1.text = "• " + bt
        r1.font.bold = True
        r1.font.size = Pt(10)
        r1.font.color.rgb = TEXT_DARK
        r2 = pi.add_run()
        r2.text = bc
        r2.font.size = Pt(10)
        r2.font.color.rgb = TEXT_MUTED

    # Box 4: storytelling/ & learnings/
    add_card(s5, Inches(6.8), Inches(4.45), gw, gh)
    tb = s5.shapes.add_textbox(Inches(7.0), Inches(4.55), gw - Inches(0.4), gh - Inches(0.2))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "storytelling/ & learnings/... (Dissemination & Evolution)"
    r.font.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = EMERALD
    s_items = [
        ("Audience Stories: ", "Executive, NOC technical, and customer-facing incident reports."),
        ("Spoken Briefs: ", "Voice-curated summaries stripped of technical syntax."),
        ("FCAPS Learnings: ", "Operator-reviewed knowledge preventing recurrence."),
    ]
    for bt, bc in s_items:
        pi = tf.add_paragraph()
        pi.space_before = Pt(6)
        r1 = pi.add_run()
        r1.text = "• " + bt
        r1.font.bold = True
        r1.font.size = Pt(10)
        r1.font.color.rgb = TEXT_DARK
        r2 = pi.add_run()
        r2.text = bc
        r2.font.size = Pt(10)
        r2.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 6: Use Case 1 - Multi-Domain Alarm Correlation
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_bg(s6, SLATE_BG)
    add_header(s6, "Use Case 1", "Inter- & Intra-Domain Alarm Correlation Engine", "Ingesting raw alerts from disparate silos and clustering them into verified incident graphs")

    add_card(s6, Inches(0.8), Inches(1.85), col_w, Inches(4.9))
    tb = s6.shapes.add_textbox(Inches(1.0), Inches(2.0), col_w - Inches(0.4), Inches(4.6))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "Correlation Engine Workflow"
    r.font.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = NAVY_DARK

    c_wf = [
        ("Multi-Domain Ingestion: ", "Ingests alarms across RAN, Transport, Mobile Core, and Site Power."),
        ("Deduplication & Normalization: ", "Removes duplicate events via stable fingerprint hash."),
        ("10-Minute Sliding Window: ", "Groups alarms exhibiting temporal co-occurrence."),
        ("Topology & Service Traversal: ", "Checks reachability and shared service procedure impact."),
        ("Intent Evaluation: ", "Violations of service objectives (e.g. 4G Attach SR) increase candidate score."),
        ("gbrain MCP Writer: ", "Upserts clusters, decisions, and hypotheses into gbrain."),
    ]
    for bt, bc in c_wf:
        pi = tf.add_paragraph()
        pi.space_before = Pt(8)
        r1 = pi.add_run()
        r1.text = "• " + bt
        r1.font.bold = True
        r1.font.size = Pt(11)
        r1.font.color.rgb = TEXT_DARK
        r2 = pi.add_run()
        r2.text = bc
        r2.font.size = Pt(11)
        r2.font.color.rgb = TEXT_MUTED

    add_card(s6, Inches(6.8), Inches(1.85), col_w, Inches(4.9))
    tb = s6.shapes.add_textbox(Inches(7.0), Inches(2.0), col_w - Inches(0.4), Inches(4.6))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "Outcome Classification & Scope"
    r.font.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = BLUE_ACCENT

    p_sub = tf.add_paragraph()
    p_sub.space_before = Pt(8)
    r_sub = p_sub.add_run()
    r_sub.text = "Three Operational Classifications:"
    r_sub.font.bold = True
    r_sub.font.size = Pt(11.5)
    r_sub.font.color.rgb = NAVY_DARK

    class_items = [
        ("Standalone Event: ", "Low severity or isolated alarm; retained for audit, no escalation.", TEXT_MUTED),
        ("Candidate Incident: ", "Watch cluster gathering telemetry evidence (status: candidate).", AMBER),
        ("Confirmed Incident: ", "Validated correlation surpassing threshold; promoted for storytelling.", RED),
    ]
    for bt, bc, col in class_items:
        pi = tf.add_paragraph()
        pi.space_before = Pt(6)
        r1 = pi.add_run()
        r1.text = "• " + bt
        r1.font.bold = True
        r1.font.size = Pt(10.5)
        r1.font.color.rgb = col
        r2 = pi.add_run()
        r2.text = bc
        r2.font.size = Pt(10.5)
        r2.font.color.rgb = TEXT_MUTED

    p_scope = tf.add_paragraph()
    p_scope.space_before = Pt(12)
    r_scope = p_scope.add_run()
    r_scope.text = "Correlation Scope Separation:"
    r_scope.font.bold = True
    r_scope.font.size = Pt(11.5)
    r_scope.font.color.rgb = NAVY_DARK

    scope_items = [
        ("Intra-Domain: ", "Contained inside one network area (e.g. 5G Core AMF-SMF or Power UPS)."),
        ("Inter-Domain: ", "Cross-boundary cascade (e.g. Transport jitter causing IMS SIP drops)."),
    ]
    for bt, bc in scope_items:
        pi = tf.add_paragraph()
        pi.space_before = Pt(6)
        r1 = pi.add_run()
        r1.text = "• " + bt
        r1.font.bold = True
        r1.font.size = Pt(10.5)
        r1.font.color.rgb = TEXT_DARK
        r2 = pi.add_run()
        r2.text = bc
        r2.font.size = Pt(10.5)
        r2.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 7: Use Case 2 - Deterministic Incident Storytelling
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    set_bg(s7, SLATE_BG)
    add_header(s7, "Use Case 2", "Deterministic Incident Storytelling (Data StoryTeller)", "Eliminating AI hallucinations by grounding all incident narratives in graph-traversed truth")

    add_card(s7, Inches(0.8), Inches(1.85), col_w, Inches(4.9))
    tb = s7.shapes.add_textbox(Inches(1.0), Inches(2.0), col_w - Inches(0.4), Inches(4.6))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "Graph-Grounded Narration"
    r.font.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = NAVY_DARK

    st_items = [
        ("Zero Hallucination: ", "The LLM never invents network functions, timestamps, or root causes."),
        ("MobileCoreKnowledge: ", "Directly queries gbrain through GbrainClient to assemble the verified fact set."),
        ("Strict Narrative Schema: ", "Standard 8-part incident structure: Executive Summary, Impact, Leading Hypothesis, Correlation, Grouping Rationale, Causal Chain, Timeline, Still Open."),
        ("Auditability: ", "Every statement links back to source alarm IDs, Prometheus metrics, or PCAPs."),
    ]
    for bt, bc in st_items:
        pi = tf.add_paragraph()
        pi.space_before = Pt(10)
        r1 = pi.add_run()
        r1.text = "• " + bt
        r1.font.bold = True
        r1.font.size = Pt(11)
        r1.font.color.rgb = TEXT_DARK
        r2 = pi.add_run()
        r2.text = bc
        r2.font.size = Pt(11)
        r2.font.color.rgb = TEXT_MUTED

    add_card(s7, Inches(6.8), Inches(1.85), col_w, Inches(4.9))
    tb = s7.shapes.add_textbox(Inches(7.0), Inches(2.0), col_w - Inches(0.4), Inches(4.6))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "Tri-Audience Tailored Briefs"
    r.font.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = BLUE_ACCENT

    aud_items = [
        ("1. Executive Brief: ", "High-level summary of subscriber impact (14.2k users), business revenue risk, SLA breach status, and recovery ETA.", PURPLE),
        ("2. NOC Technical Brief: ", "Deep technical diagnosis: node names (AMF-01, PGW-01), specific alarm codes (CPU_HIGH, S1AP_DROP), 3GPP causes (NAS Cause 22), and exact diagnostic steps.", NAVY_DARK),
        ("3. Customer Operations Update: ", "Transparent, jargon-free notifications for enterprise clients and helpdesk teams explaining regional connectivity impact.", EMERALD),
    ]
    for bt, bc, col in aud_items:
        pi = tf.add_paragraph()
        pi.space_before = Pt(12)
        r1 = pi.add_run()
        r1.text = bt
        r1.font.bold = True
        r1.font.size = Pt(11)
        r1.font.color.rgb = col
        r2 = pi.add_run()
        r2.text = bc
        r2.font.size = Pt(11)
        r2.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 8: Production Golden Scenarios Implemented
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    set_bg(s8, SLATE_BG)
    add_header(s8, "Implemented Scenarios", "Production Golden Incident Scenarios in gbrain", "Four end-to-end multi-domain incident archetypes implemented, validated, and regression-tested")

    # 4 Cards
    add_card(s8, Inches(0.8), Inches(1.85), gw, gh)
    tb = s8.shapes.add_textbox(Inches(1.0), Inches(1.95), gw - Inches(0.4), gh - Inches(0.2))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "1. 5G Core AMF-01 Signaling Overload"
    r.font.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = RED
    amf_items = [
        ("Impact: ", "5G Initial UE Registration degradation across Dubai Core DC-1."),
        ("Causal Chain: ", "Signaling burst -> AMF CPU @ 98% -> NAS buffer drops (Cause 22)."),
        ("Remediation & Recovery: ", "Scaled AMF worker pods from 4 to 8; RSR recovered from 94.7% to 99.9%."),
    ]
    for bt, bc in amf_items:
        pi = tf.add_paragraph()
        pi.space_before = Pt(5)
        r1 = pi.add_run()
        r1.text = "• " + bt
        r1.font.bold = True
        r1.font.size = Pt(10)
        r1.font.color.rgb = TEXT_DARK
        r2 = pi.add_run()
        r2.text = bc
        r2.font.size = Pt(10)
        r2.font.color.rgb = TEXT_MUTED

    add_card(s8, Inches(6.8), Inches(1.85), gw, gh)
    tb = s8.shapes.add_textbox(Inches(7.0), Inches(1.95), gw - Inches(0.4), gh - Inches(0.2))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "2. LTE Attach & S1-MME Path Degradation"
    r.font.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = AMBER
    lte_items = [
        ("Impact: ", "14,200 mobile subscribers across Dubai North eNodeBs."),
        ("Multi-Node Telemetry: ", "MME-01 attach rejects + HSS-01 Diameter ULR timeouts + eNodeB drops."),
        ("Leading Hypothesis: ", "S1-MME transport path degradation causing control-plane packet loss."),
    ]
    for bt, bc in lte_items:
        pi = tf.add_paragraph()
        pi.space_before = Pt(5)
        r1 = pi.add_run()
        r1.text = "• " + bt
        r1.font.bold = True
        r1.font.size = Pt(10)
        r1.font.color.rgb = TEXT_DARK
        r2 = pi.add_run()
        r2.text = bc
        r2.font.size = Pt(10)
        r2.font.color.rgb = TEXT_MUTED

    add_card(s8, Inches(0.8), Inches(4.45), gw, gh)
    tb = s8.shapes.add_textbox(Inches(1.0), Inches(4.55), gw - Inches(0.4), gh - Inches(0.2))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "3. Cross-Domain IMS Voice Call Setup (CSSR)"
    r.font.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = PURPLE
    ims_items = [
        ("Cross-Boundary Cascade: ", "Correlates RAN (eNodeB-22), IMS (P-CSCF), and Backhaul."),
        ("Correlated Alarms: ", "RRC Setup Failures + SIP INVITE drop spike + Link Jitter."),
        ("Intent State: ", "Service intent `voice_call_drop` violated; correlation score: 92."),
    ]
    for bt, bc in ims_items:
        pi = tf.add_paragraph()
        pi.space_before = Pt(5)
        r1 = pi.add_run()
        r1.text = "• " + bt
        r1.font.bold = True
        r1.font.size = Pt(10)
        r1.font.color.rgb = TEXT_DARK
        r2 = pi.add_run()
        r2.text = bc
        r2.font.size = Pt(10)
        r2.font.color.rgb = TEXT_MUTED

    add_card(s8, Inches(6.8), Inches(4.45), gw, gh)
    tb = s8.shapes.add_textbox(Inches(7.0), Inches(4.55), gw - Inches(0.4), gh - Inches(0.2))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "4. Gi-LAN Bottleneck & SGi Data Disruption"
    r.font.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = BLUE_ACCENT
    sgi_items = [
        ("User-Plane Outage: ", "Mobile broadband packet forwarding collapse."),
        ("Evidence Chain: ", "PGW-01 throughput drop + NAT firewall session exhaustion + router discards."),
        ("RCA Diagnosis: ", "Gi-LAN perimeter NAT table saturation, not 3GPP Core failure."),
    ]
    for bt, bc in sgi_items:
        pi = tf.add_paragraph()
        pi.space_before = Pt(5)
        r1 = pi.add_run()
        r1.text = "• " + bt
        r1.font.bold = True
        r1.font.size = Pt(10)
        r1.font.color.rgb = TEXT_DARK
        r2 = pi.add_run()
        r2.text = bc
        r2.font.size = Pt(10)
        r2.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 9: Use Case 3 - Visual Explanation & Investigation Workspace
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    set_bg(s9, SLATE_BG)
    add_header(s9, "Use Case 3", "Visual Explanation & Interactive Investigation Workspace", "Translating complex graph relationships into actionable visual widgets for NOC operators")

    add_card(s9, Inches(0.8), Inches(1.85), col_w, Inches(4.9))
    tb = s9.shapes.add_textbox(Inches(1.0), Inches(2.0), col_w - Inches(0.4), Inches(4.6))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "Visual Explanation Service"
    r.font.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = NAVY_DARK

    ve_items = [
        ("Dedicated Visual Engine: ", "`VisualExplanationService` translates `gbrain` subgraphs into standardized frontend JSON contracts."),
        ("Inline Story Visuals: ", "Embedded directly inside chat response cards alongside text."),
        ("Expandable Evidence Drawer: ", "One-click inspection of supporting raw Prometheus/Loki claims without cluttering the chat view."),
        ("HUD Highlight Orb: ", "Visual pulse indicating active incident severity in the MARK HUD."),
    ]
    for bt, bc in ve_items:
        pi = tf.add_paragraph()
        pi.space_before = Pt(10)
        r1 = pi.add_run()
        r1.text = "• " + bt
        r1.font.bold = True
        r1.font.size = Pt(11)
        r1.font.color.rgb = TEXT_DARK
        r2 = pi.add_run()
        r2.text = bc
        r2.font.size = Pt(11)
        r2.font.color.rgb = TEXT_MUTED

    add_card(s9, Inches(6.8), Inches(1.85), col_w, Inches(4.9))
    tb = s9.shapes.add_textbox(Inches(7.0), Inches(2.0), col_w - Inches(0.4), Inches(4.6))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "Interactive NOC Investigation Canvas"
    r.font.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = BLUE_ACCENT

    canv_items = [
        ("Topology Blast Radius: ", "Interactive node graph highlighting root node, cascade path, and unaffected neighboring nodes."),
        ("Side-by-Side Telemetry: ", "Synchronized time-series metrics from Grafana pinned next to graph nodes."),
        ("Operator Action Strip: ", "Dynamic action buttons (Acknowledge, Assign, Suppress, Merge, Scale Out) bound to the selected incident ID."),
        ("Full Audit Trail: ", "Records every operator interaction and state change in Incident Registry."),
    ]
    for bt, bc in canv_items:
        pi = tf.add_paragraph()
        pi.space_before = Pt(10)
        r1 = pi.add_run()
        r1.text = "• " + bt
        r1.font.bold = True
        r1.font.size = Pt(11)
        r1.font.color.rgb = TEXT_DARK
        r2 = pi.add_run()
        r2.text = bc
        r2.font.size = Pt(11)
        r2.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 10: Use Case 4 - Mark Conversational Assistant & Voice
    # =========================================================================
    s10 = prs.slides.add_slide(blank_layout)
    set_bg(s10, SLATE_BG)
    add_header(s10, "Use Case 4", "Mark Conversational Assistant & Curated Voice", "Hands-free operations in the NOC: low-latency speech synthesis tailored for operational dialogue")

    # Banner
    add_card(s10, Inches(0.8), Inches(1.85), Inches(11.7), Inches(0.95), bg_color=RGBColor(245, 243, 255), border_color=RGBColor(221, 214, 254))
    tb_vban = s10.shapes.add_textbox(Inches(1.0), Inches(1.9), Inches(11.3), Inches(0.85))
    tf_vban = tb_vban.text_frame
    tf_vban.word_wrap = True
    p = tf_vban.paragraphs[0]
    r1 = p.add_run()
    r1.text = "Core Voice Rule: "
    r1.font.bold = True
    r1.font.size = Pt(13)
    r1.font.color.rgb = PURPLE
    r2 = p.add_run()
    r2.text = "\"Brain content is context, NOT a script to read aloud.\""
    r2.font.bold = True
    r2.font.size = Pt(13)
    r2.font.color.rgb = NAVY_DARK

    p_sub = tf_vban.add_paragraph()
    p_sub.space_before = Pt(3)
    r3 = p_sub.add_run()
    r3.text = "Mark never reads raw markdown, incident slugs, UUIDs, punctuation, bullet characters, or table syntax over voice. Technical details stay on screen; speech is conversational and brief."
    r3.font.size = Pt(10.5)
    r3.font.color.rgb = TEXT_MUTED

    # Left: Pipeline
    add_card(s10, Inches(0.8), col_y, col_w, col_h)
    tb = s10.shapes.add_textbox(Inches(1.0), col_y + Inches(0.15), col_w - Inches(0.4), col_h - Inches(0.3))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "Voice Curation Pipeline"
    r.font.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = NAVY_DARK

    vp_items = [
        ("Dual-Payload Generation: ", "Every turn outputs a rich technical written report and a separate spoken brief."),
        ("VAD Noise Rejection: ", "Silero VAD thresholding (0.70) ignores noisy NOC background chatter."),
        ("Streaming WebSocket: ", "Sub-400ms audio delivery via ElevenLabs streaming synthesis."),
    ]
    for bt, bc in vp_items:
        pi = tf.add_paragraph()
        pi.space_before = Pt(12)
        r1 = pi.add_run()
        r1.text = "• " + bt
        r1.font.bold = True
        r1.font.size = Pt(11)
        r1.font.color.rgb = TEXT_DARK
        r2 = pi.add_run()
        r2.text = bc
        r2.font.size = Pt(11)
        r2.font.color.rgb = TEXT_MUTED

    # Right: Example
    add_card(s10, Inches(6.8), col_y, col_w, col_h)
    tb = s10.shapes.add_textbox(Inches(7.0), col_y + Inches(0.15), col_w - Inches(0.4), col_h - Inches(0.3))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "Example: Written vs. Spoken Output"
    r.font.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = PURPLE

    p_w1 = tf.add_paragraph()
    p_w1.space_before = Pt(8)
    r = p_w1.add_run()
    r.text = "On-Screen Written Report (Full Technical):"
    r.font.bold = True
    r.font.size = Pt(10)
    r.font.color.rgb = TEXT_MUTED

    p_w2 = tf.add_paragraph()
    p_w2.space_before = Pt(2)
    r = p_w2.add_run()
    r.text = "# Incident story — incidents/mobile-core/amf-overload-2026-08-09\nStatus: resolved | Severity: SEV-2 | Score: 0.94\nRoot cause: AMF-01 CPU saturation (98%) on worker process 104.\nRemediation: Pod replica count increased 4 -> 8."
    r.font.size = Pt(9.5)
    r.font.color.rgb = NAVY_DARK

    p_s1 = tf.add_paragraph()
    p_s1.space_before = Pt(10)
    r = p_s1.add_run()
    r.text = "Curated Spoken Voice Brief (Audio Output):"
    r.font.bold = True
    r.font.size = Pt(10)
    r.font.color.rgb = PURPLE

    p_s2 = tf.add_paragraph()
    p_s2.space_before = Pt(2)
    r = p_s2.add_run()
    r.text = "\"UE Registration is currently severity two and resolved. The confirmed root cause was AMF-01 CPU saturation from a registration signaling burst. Capacity was scaled out and service recovered.\""
    r.font.size = Pt(10.5)
    r.font.italic = True
    r.font.color.rgb = NAVY_DARK

    # =========================================================================
    # SLIDE 11: Use Cases 5 & 6 - RCA Triage & Telemetry Projection
    # =========================================================================
    s11 = prs.slides.add_slide(blank_layout)
    set_bg(s11, SLATE_BG)
    add_header(s11, "Use Cases 5 & 6", "RCA Triage Advisory & Live Telemetry Projection", "Hypothesis-driven root cause analysis coupled with high-speed Grafana LGTM data projection")

    add_card(s11, Inches(0.8), Inches(1.85), col_w, Inches(4.9))
    tb = s11.shapes.add_textbox(Inches(1.0), Inches(2.0), col_w - Inches(0.4), Inches(4.6))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "Use Case 5: Root Cause Analysis (RCA)"
    r.font.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = NAVY_DARK

    rca_items = [
        ("Hypothesis Engine: ", "Evaluates candidate causal explanations against confirmed graph evidence."),
        ("Confidence Scoring: ", "Scores hypotheses (0.0 to 1.0) based on metric alignment, topology, and alarms."),
        ("Fact vs. Inference: ", "Explicitly separates confirmed facts from working inferences."),
        ("Diagnostic Next Steps: ", "Suggests the exact next test or command needed to confirm unproven hypotheses."),
        ("Remediation Advisory: ", "Proposes approved MOPs (Method of Procedure) with human approval gates."),
    ]
    for bt, bc in rca_items:
        pi = tf.add_paragraph()
        pi.space_before = Pt(9)
        r1 = pi.add_run()
        r1.text = "• " + bt
        r1.font.bold = True
        r1.font.size = Pt(10.5)
        r1.font.color.rgb = TEXT_DARK
        r2 = pi.add_run()
        r2.text = bc
        r2.font.size = Pt(10.5)
        r2.font.color.rgb = TEXT_MUTED

    add_card(s11, Inches(6.8), Inches(1.85), col_w, Inches(4.9))
    tb = s11.shapes.add_textbox(Inches(7.0), Inches(2.0), col_w - Inches(0.4), Inches(4.6))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "Use Case 6: Telemetry Evidence Projection"
    r.font.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = BLUE_ACCENT

    tel_items = [
        ("Grafana MCP Integration: ", "Connects Mark directly to Prometheus metrics, Loki logs, and Tempo traces via MCP."),
        ("Compact Projections: ", "Projects only salient threshold breaches into gbrain (`grafana/*` pages), avoiding database bloat."),
        ("On-Demand Enrichment: ", "Storyteller queries live Grafana telemetry on demand when `include_live_telemetry` is requested."),
        ("Bidirectional Provenance: ", "Graph claims embed direct deep-links to Grafana dashboard time windows."),
    ]
    for bt, bc in tel_items:
        pi = tf.add_paragraph()
        pi.space_before = Pt(9)
        r1 = pi.add_run()
        r1.text = "• " + bt
        r1.font.bold = True
        r1.font.size = Pt(10.5)
        r1.font.color.rgb = TEXT_DARK
        r2 = pi.add_run()
        r2.text = bc
        r2.font.size = Pt(10.5)
        r2.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 12: Use Case 7 - FCAPS Closed-Loop Learning
    # =========================================================================
    s12 = prs.slides.add_slide(blank_layout)
    set_bg(s12, SLATE_BG)
    add_header(s12, "Use Case 7", "FCAPS Closed-Loop Learning & Knowledge Evolution", "Transforming resolved operational incidents into durable enterprise knowledge and runbooks")

    c_w = Inches(3.64)
    c_h = Inches(4.9)

    add_card(s12, Inches(0.8), Inches(1.85), c_w, c_h)
    tb = s12.shapes.add_textbox(Inches(0.95), Inches(2.0), c_w - Inches(0.3), c_h - Inches(0.3))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "The FCAPS Operational Lens"
    r.font.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = NAVY_DARK
    fcaps_items = [
        ("F - Fault: ", "Alarms, link down, crash loops, timeouts."),
        ("C - Configuration: ", "Change requests, MOP execution, rollbacks."),
        ("A - Accounting: ", "Billing, charging gateway, quota limits."),
        ("P - Performance: ", "Throughput, latency, CSSR, CPU saturation."),
        ("S - Security: ", "Auth anomalies, certificate expiry, DDoS."),
    ]
    for bt, bc in fcaps_items:
        pi = tf.add_paragraph()
        pi.space_before = Pt(8)
        r1 = pi.add_run()
        r1.text = "• " + bt
        r1.font.bold = True
        r1.font.size = Pt(10.5)
        r1.font.color.rgb = TEXT_DARK
        r2 = pi.add_run()
        r2.text = bc
        r2.font.size = Pt(10.5)
        r2.font.color.rgb = TEXT_MUTED

    add_card(s12, Inches(4.84), Inches(1.85), c_w, c_h)
    tb = s12.shapes.add_textbox(Inches(4.99), Inches(2.0), c_w - Inches(0.3), c_h - Inches(0.3))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "Post-Incident Enrichment"
    r.font.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = BLUE_ACCENT
    enr_items = [
        ("Operator Feedback Loop: ", "NOC operators review and validate incident resolutions."),
        ("Gap Identification: ", "Detects missing telemetry metrics or incomplete alarm definitions."),
        ("Runbook Candidates: ", "Automatically proposes updated diagnostic runbooks in `learnings/`."),
        ("Review Gates: ", "Learnings require senior engineer signoff before becoming permanent knowledge."),
    ]
    for bt, bc in enr_items:
        pi = tf.add_paragraph()
        pi.space_before = Pt(8)
        r1 = pi.add_run()
        r1.text = "• " + bt
        r1.font.bold = True
        r1.font.size = Pt(10.5)
        r1.font.color.rgb = TEXT_DARK
        r2 = pi.add_run()
        r2.text = bc
        r2.font.size = Pt(10.5)
        r2.font.color.rgb = TEXT_MUTED

    add_card(s12, Inches(8.88), Inches(1.85), c_w, c_h)
    tb = s12.shapes.add_textbox(Inches(9.03), Inches(2.0), c_w - Inches(0.3), c_h - Inches(0.3))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "Durable Knowledge Evolution"
    r.font.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = EMERALD
    evo_items = [
        ("Preventing Recurrence: ", "Future correlation runs match known past incident signatures instantly."),
        ("Query Improvement: ", "Generates optimized PromQL/LogQL queries for emerging failure modes."),
        ("Asset Health Index: ", "Continuously refines component MTBF and risk profiles."),
        ("Autonomous Improvement: ", "System evolves with every shift without software code changes."),
    ]
    for bt, bc in evo_items:
        pi = tf.add_paragraph()
        pi.space_before = Pt(8)
        r1 = pi.add_run()
        r1.text = "• " + bt
        r1.font.bold = True
        r1.font.size = Pt(10.5)
        r1.font.color.rgb = TEXT_DARK
        r2 = pi.add_run()
        r2.text = bc
        r2.font.size = Pt(10.5)
        r2.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 13: Capability Matrix
    # =========================================================================
    s13 = prs.slides.add_slide(blank_layout)
    set_bg(s13, SLATE_BG)
    add_header(s13, "Capability Matrix", "Summary of Implemented gbrain Capabilities & Services", "Comprehensive overview of services, connectors, data consumed, and operational outputs")

    add_card(s13, Inches(0.8), Inches(1.85), Inches(11.7), Inches(4.9))
    tb = s13.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.3), Inches(4.6))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "Implemented Service Capabilities in KAgent / MARK"
    r.font.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = NAVY_DARK

    svc_matrix = [
        ("Correlation Service: ", "Groups alarms/KPIs via time-window, topology, and service intent. Writes correlation clusters & decisions to gbrain.", BLUE_ACCENT),
        ("Storytelling Service: ", "Traverses gbrain graph to produce deterministic written stories, executive briefs, customer updates, and spoken audio.", PURPLE),
        ("Visual Explanation Service: ", "Converts gbrain incident graphs into visual widgets, HUD highlight orbs, evidence drawers, and mini-canvas investigation views.", EMERALD),
        ("Telemetry Evidence Service: ", "Bridges Grafana LGTM with gbrain via MCP Hub, extracting facts and maintaining bidirectional deep-link provenance.", AMBER),
        ("RCA & Triage Service: ", "Generates scored causal hypotheses, differentiates facts from inferences, and formulates diagnostic next steps.", RED),
        ("FCAPS Learning Service: ", "Applies 5-dimension operational lens to capture reviewed learnings, playbook candidates, and query templates in gbrain.", CYAN_ACCENT),
    ]
    for bt, bc, col in svc_matrix:
        pi = tf.add_paragraph()
        pi.space_before = Pt(10)
        r1 = pi.add_run()
        r1.text = "• " + bt
        r1.font.bold = True
        r1.font.size = Pt(11)
        r1.font.color.rgb = col
        r2 = pi.add_run()
        r2.text = bc
        r2.font.size = Pt(11)
        r2.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 14: Business Value & Operational Impact
    # =========================================================================
    s14 = prs.slides.add_slide(blank_layout)
    set_bg(s14, SLATE_BG)
    add_header(s14, "Operational Impact", "Quantified Business Value & ROI for Telecom Operators", "Realizing dramatic MTTR reductions, alert fatigue elimination, and total reporting auditability")

    m_w = Inches(2.7)
    m_h = Inches(4.9)

    # M1: MTTR
    add_card(s14, Inches(0.8), Inches(1.85), m_w, m_h)
    tb = s14.shapes.add_textbox(Inches(0.95), Inches(2.0), m_w - Inches(0.3), m_h - Inches(0.3))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "-75%"
    r.font.bold = True
    r.font.size = Pt(36)
    r.font.color.rgb = BLUE_ACCENT
    p2 = tf.add_paragraph()
    r = p2.add_run()
    r.text = "MTTR Reduction"
    r.font.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = NAVY_DARK
    m1_items = [
        "Triages root causes in seconds instead of hours.",
        "Automates multi-domain correlation across Core, RAN, and Transport.",
        "Immediate diagnostic guidance for tier-1 NOC engineers.",
    ]
    for it in m1_items:
        pi = tf.add_paragraph()
        pi.space_before = Pt(8)
        r = pi.add_run()
        r.text = "• " + it
        r.font.size = Pt(10)
        r.font.color.rgb = TEXT_MUTED

    # M2: Noise
    add_card(s14, Inches(3.8), Inches(1.85), m_w, m_h)
    tb = s14.shapes.add_textbox(Inches(3.95), Inches(2.0), m_w - Inches(0.3), m_h - Inches(0.3))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "-90%"
    r.font.bold = True
    r.font.size = Pt(36)
    r.font.color.rgb = EMERALD
    p2 = tf.add_paragraph()
    r = p2.add_run()
    r.text = "Alert Fatigue"
    r.font.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = NAVY_DARK
    m2_items = [
        "Suppresses redundant symptom alarms into single incidents.",
        "Separates standalone noise from critical promotions.",
        "Prevents alarm storms during fiber cuts or power outages.",
    ]
    for it in m2_items:
        pi = tf.add_paragraph()
        pi.space_before = Pt(8)
        r = pi.add_run()
        r.text = "• " + it
        r.font.size = Pt(10)
        r.font.color.rgb = TEXT_MUTED

    # M3: Trust
    add_card(s14, Inches(6.8), Inches(1.85), m_w, m_h)
    tb = s14.shapes.add_textbox(Inches(6.95), Inches(2.0), m_w - Inches(0.3), m_h - Inches(0.3))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "100%"
    r.font.bold = True
    r.font.size = Pt(36)
    r.font.color.rgb = PURPLE
    p2 = tf.add_paragraph()
    r = p2.add_run()
    r.text = "Deterministic Trust"
    r.font.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = NAVY_DARK
    m3_items = [
        "Zero AI hallucination in mission-critical NOC narratives.",
        "Complete audit trail from claim to Prometheus/Loki records.",
        "Compliant with telecom regulatory and SLA audit mandates.",
    ]
    for it in m3_items:
        pi = tf.add_paragraph()
        pi.space_before = Pt(8)
        r = pi.add_run()
        r.text = "• " + it
        r.font.size = Pt(10)
        r.font.color.rgb = TEXT_MUTED

    # M4: Efficiency
    add_card(s14, Inches(9.8), Inches(1.85), m_w, m_h)
    tb = s14.shapes.add_textbox(Inches(9.95), Inches(2.0), m_w - Inches(0.3), m_h - Inches(0.3))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "3x"
    r.font.bold = True
    r.font.size = Pt(36)
    r.font.color.rgb = AMBER
    p2 = tf.add_paragraph()
    r = p2.add_run()
    r.text = "NOC Efficiency"
    r.font.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = NAVY_DARK
    m4_items = [
        "Hands-free voice querying for incident commanders.",
        "Automated executive and customer notification generation.",
        "Faster shift handovers with persistent incident storytelling.",
    ]
    for it in m4_items:
        pi = tf.add_paragraph()
        pi.space_before = Pt(8)
        r = pi.add_run()
        r.text = "• " + it
        r.font.size = Pt(10)
        r.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 15: Conclusion & Future Roadmap (Dark Theme)
    # =========================================================================
    s15 = prs.slides.add_slide(blank_layout)
    set_bg(s15, NAVY_DARK)
    card15 = add_card(s15, Inches(0.8), Inches(0.8), Inches(11.73), Inches(5.9), bg_color=NAVY_CARD, border_color=BORDER_DARK)

    tb15 = s15.shapes.add_textbox(Inches(1.4), Inches(1.3), Inches(10.5), Inches(5.0))
    tf15 = tb15.text_frame
    tf15.word_wrap = True

    p = tf15.paragraphs[0]
    r = p.add_run()
    r.text = "The Future of Autonomous Telecom with gbrain"
    r.font.name = "Segoe UI"
    r.font.size = Pt(28)
    r.font.bold = True
    r.font.color.rgb = WHITE

    p_sub = tf15.add_paragraph()
    p_sub.space_before = Pt(8)
    r = p_sub.add_run()
    r.text = "Next Development Horizons on the KAgent Platform"
    r.font.name = "Segoe UI"
    r.font.size = Pt(15)
    r.font.color.rgb = CYAN_ACCENT

    road_items = [
        ("Autonomous Closed-Loop Remediation: ", "Linking gbrain root cause decisions to approved Automation Engine playbooks with canary rollback safety."),
        ("Audio-Visual Synchronized Playback: ", "Claim-level visual highlight synchronization with streaming spoken audio briefs in Mark HUD."),
        ("Live Digital Twin Ingestion: ", "Real-time topology synchronization streaming from network discovery engines into gbrain."),
        ("Multi-Operator Federated Knowledge: ", "Cross-network anomaly pattern exchange without leaking confidential subscriber metadata."),
    ]
    for bt, bc in road_items:
        pi = tf15.add_paragraph()
        pi.space_before = Pt(14)
        r1 = pi.add_run()
        r1.text = "• " + bt
        r1.font.name = "Segoe UI"
        r1.font.bold = True
        r1.font.size = Pt(12)
        r1.font.color.rgb = WHITE
        r2 = pi.add_run()
        r2.text = bc
        r2.font.name = "Segoe UI"
        r2.font.size = Pt(12)
        r2.font.color.rgb = TEXT_LIGHT

    p_end = tf15.add_paragraph()
    p_end.space_before = Pt(24)
    r_end = p_end.add_run()
    r_end.text = "Thank You  |  Q&A  |  adeel@kagent.dev"
    r_end.font.name = "Segoe UI"
    r_end.font.size = Pt(13)
    r_end.font.bold = True
    r_end.font.color.rgb = EMERALD


def main():
    out_path = "/Users/adeelarshad/kagent/gbrain_use_cases_presentation.pptx"
    if len(sys.argv) > 1:
        out_path = sys.argv[1]

    prs = Presentation()
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT

    build_all_slides(prs)
    prs.save(out_path)
    print(f"Successfully generated standard PPTX: {out_path} (15 slides)")


if __name__ == "__main__":
    main()
